"""The Excel loader's two workbook-failure log lines carry no cell value.

`load_excel_sheets` logs when the sheet-visibility probe (openpyxl) or the
shared calamine handle fails to open. Both libraries quote what they choked
on — a cell, a shared string, a sheet XML fragment — in the exception text,
and the workbook is a CUSTOMER upload. Article IV's value rule applies: such
a line carries the exception TYPE and a fixed reason, never `{e}`. Escaping
alone would still put the customer's value into the log, and a newline in it
would forge an extra record in the newline-delimited client log.

Both failures are forced with fakes (no real workbook is needed), and the
log is captured at `logger_utils.log_with_sid`, which the loader imports
inside the function.
"""
import pandas as pd
import pytest

import excel_table_detector
import logger_utils

SENTINEL = "CELLVALUE_4711_secret_salary"
FORGED = "\n2026-09-24 00:00:00,000 | INFO | [sid=x] EXEC_OK job_id=forged"


class _CellQuotingError(ValueError):
    """Stands in for a library error that quotes the offending cell."""


def _hostile_error():
    return _CellQuotingError(f"cannot parse cell value '{SENTINEL}'{FORGED}")


class _BrokenSheet:
    @property
    def sheet_state(self):
        raise _hostile_error()


class _FakeWorkbook:
    sheetnames = ["Sheet1"]

    def __getitem__(self, name):
        return _BrokenSheet()

    def close(self):
        pass


@pytest.fixture
def captured(monkeypatch, tmp_path):
    import openpyxl

    lines = []

    def record(sid, level, message, *args, **kwargs):
        lines.append({"sid": sid, "level": level, "message": str(message),
                      "kwargs": {k: str(v) for k, v in kwargs.items()}})

    monkeypatch.setattr(logger_utils, "log_with_sid", record)
    monkeypatch.setattr(openpyxl, "load_workbook", lambda *a, **k: _FakeWorkbook())

    def broken_excel_file(*args, **kwargs):
        raise _hostile_error()

    monkeypatch.setattr(pd, "ExcelFile", broken_excel_file)
    # The per-sheet pipeline is not under test; it must not touch the fake.
    monkeypatch.setattr(excel_table_detector, "_detect_and_extract_table",
                        lambda *a, **k: None)
    path = tmp_path / "book.xlsx"
    path.write_bytes(b"not really a workbook")
    excel_table_detector.load_excel_sheets(path, "book.xlsx")
    return lines


def _rendered(row) -> str:
    return " ".join([row["message"]] + [f"{k}={v}" for k, v in row["kwargs"].items()])


@pytest.mark.parametrize("event", ["EXCEL_SHEET_STATE_PROBE_FAILED",
                                   "EXCEL_SHARED_HANDLE_FAILED"])
def test_the_failure_line_names_the_type_and_not_the_value(captured, event):
    hits = [row for row in captured if event in row["message"]]
    assert len(hits) == 1, [row["message"] for row in captured]
    line = _rendered(hits[0])
    assert "_CellQuotingError" in line, line
    assert SENTINEL not in line, line
    assert "\n" not in line and "\r" not in line, repr(line)
    assert "EXEC_OK" not in line, line


def test_no_captured_line_repeats_the_cell_value(captured):
    leaked = [row["message"] for row in captured if SENTINEL in _rendered(row)]
    assert leaked == [], leaked
